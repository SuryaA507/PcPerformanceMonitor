from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "athi rpa final - sample.docx"
OUTPUT = ROOT / "PC_Performance_Monitor_Bot_Project_Report.docx"


def clear_document(doc):
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def paragraph(doc, text="", style=None, align=None, bold=False, size=None, before=None, after=None):
    paragraph_styles = [s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH]
    if style and style not in paragraph_styles:
        style = None
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def page_break(doc):
    doc.add_page_break()


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_body(doc, text):
    paragraph(doc, text, style="Body Text", after=6)


def add_heading(doc, text, level=1):
    p = paragraph(doc, text, style=f"Heading {level}", bold=True, after=8)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def build_report():
    doc = Document(TEMPLATE)
    clear_document(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Cover page
    paragraph(doc, "PC PERFORMANCE MONITOR BOT", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, after=24)
    paragraph(doc, "A PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=4)
    paragraph(doc, "on", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    paragraph(doc, "ROBOTIC PROCESS AUTOMATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=28)
    paragraph(doc, "Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    paragraph(doc, "ATHITYAA A", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=22)
    paragraph(doc, "in partial fulfillment of the requirement for the award of the degree", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "of", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "BACHELOR OF TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    paragraph(doc, "in", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "ARTIFICIAL INTELLIGENCE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    paragraph(doc, "AND DATA SCIENCE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=20)
    paragraph(doc, "K.S. RANGASAMY COLLEGE OF TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    paragraph(doc, "(Autonomous)", align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, "TIRUCHENGODE - 637 215", align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
    paragraph(doc, "MAY 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    page_break(doc)

    # Certificate
    paragraph(doc, "K.S. RANGASAMY COLLEGE OF TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12)
    paragraph(doc, "TIRUCHENGODE - 637 215", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, after=18)
    paragraph(doc, "BONAFIDE CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, after=22)
    add_body(doc, 'Certified that this project report titled "PC Performance Monitor Bot" is the bonafide work of ATHITYAA A who carried out the project work under my supervision. The project demonstrates the use of Robotic Process Automation to monitor real-time computer performance and present the collected data through a browser-based dashboard.')
    paragraph(doc, "", after=26)
    paragraph(doc, "SIGNATURE", bold=True)
    paragraph(doc, "Dr. S. SARUMATHI, M.E., Ph.D.,", bold=True)
    paragraph(doc, "HEAD OF THE DEPARTMENT")
    paragraph(doc, "Professor")
    paragraph(doc, "Department of Artificial Intelligence and Data Science")
    paragraph(doc, "K.S. Rangasamy College of Technology")
    paragraph(doc, "Tiruchengode - 637 215", after=24)
    paragraph(doc, "Submitted for the viva-voce examination held on", style="Body Text", after=40)
    paragraph(doc, "Internal Examiner 1                                      Internal Examiner 2", align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # Declaration
    paragraph(doc, "DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, after=22)
    add_body(doc, 'I declare that the project report on "PC PERFORMANCE MONITOR BOT" is the result of original work done by me to the best of my knowledge. This project report has not been submitted to any other university or institution for the award of any degree or diploma.')
    paragraph(doc, "", after=40)
    paragraph(doc, "Signature", align=WD_ALIGN_PARAGRAPH.RIGHT, after=30)
    paragraph(doc, "ATHITYAA A", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, after=48)
    paragraph(doc, "Place: Tiruchengode")
    paragraph(doc, "Date:")
    page_break(doc)

    add_heading(doc, "ACKNOWLEDGEMENT")
    add_body(doc, "I wish to express my sincere gratitude to our honourable Chairman Thiru. R. SRINIVASAN, BBM., MISTE., for providing immense facilities at our institution.")
    add_body(doc, "I would like to express my sincere gratitude to Principal Dr. R. GOPALAKRISHNAN, M.E., Ph.D., for providing the necessary resources and facilities for the completion of this project.")
    add_body(doc, "I proudly render my immense gratitude to the Head of the Department Dr. S. SARUMATHI, M.E., Ph.D., for her guidance and support in shaping the direction of this project.")
    add_body(doc, "I extend my sincere thanks to all faculty members of the Artificial Intelligence and Data Science Department for their valuable suggestions, cooperation, and encouragement throughout the development of this project.")
    add_body(doc, "I also acknowledge the support received from my friends and well-wishers during the preparation, testing, and documentation stages of this project.")
    page_break(doc)

    add_heading(doc, "ABSTRACT")
    add_body(doc, "The PC Performance Monitor Bot is an RPA-based system designed to monitor live computer performance and display the information through a modern web dashboard. The backend is developed using UiPath and PowerShell, while the frontend is built using HTML, CSS, and JavaScript. The system collects CPU usage, RAM usage, GPU usage, top CPU-consuming applications, top RAM-consuming applications, and applications using GPU resources.")
    add_body(doc, "The collected data is written into a JSON file named performance.json inside the dashboard folder. The browser dashboard fetches this JSON file every three seconds and updates progress bars and tables automatically. This architecture avoids complex database setup and provides a simple, beginner-friendly communication method between UiPath and the frontend.")
    add_body(doc, "The project also includes CSV and Excel export features, allowing the monitored data to be downloaded for reporting and analysis. The system is useful for students, developers, and administrators who need a lightweight real-time monitoring tool similar to Task Manager.")
    page_break(doc)

    add_heading(doc, "TABLE OF CONTENT")
    toc_rows = [
        ("CHAPTER 1", "Introduction", "1"),
        ("CHAPTER 2", "Literature Review", "4"),
        ("CHAPTER 3", "Methodology and System Implementation", "8"),
        ("CHAPTER 4", "Results and Discussion", "13"),
        ("CHAPTER 5", "Conclusion and Future Work", "18"),
        ("", "References", "20"),
    ]
    add_table(doc, ["Chapter", "Title", "Page No."], toc_rows)
    page_break(doc)

    add_heading(doc, "CHAPTER 1 INTRODUCTION")
    add_heading(doc, "ABOUT THE PROJECT")
    add_body(doc, "PC Performance Monitor Bot is a real-time monitoring project developed using UiPath as the automation backend and a browser-based dashboard as the user interface. The goal of the project is to show important system resource details such as CPU, RAM, and GPU usage in a simple and readable form.")
    add_body(doc, "The bot continuously collects performance information from Windows using PowerShell commands invoked through UiPath. The output is converted into structured JSON and stored in the dashboard folder. The frontend reads this JSON file using JavaScript fetch logic and refreshes the displayed values automatically.")
    add_heading(doc, "PROJECT SCOPE")
    add_body(doc, "The scope of this project includes live monitoring of overall CPU usage, memory usage, GPU usage, top CPU-consuming processes, top memory-consuming processes, and active GPU-using processes. The dashboard provides real-time visual feedback through progress bars, tables, status messages, and export options.")
    add_body(doc, "The project is intended for local machine monitoring and educational RPA demonstration. It shows how UiPath can interact with operating system counters, how JSON can be used as a lightweight data exchange format, and how a frontend dashboard can visualize automation output.")
    add_heading(doc, "UIPATH OBJECTIVE", 2)
    add_body(doc, "The primary UiPath objective is to automate repeated performance data collection without manual intervention. UiPath runs the backend workflow continuously, invokes PowerShell, manages repeated execution every few seconds, and ensures that the JSON file is updated for the dashboard.")
    add_heading(doc, "AIM AND SCOPE")
    add_body(doc, "The aim of the project is to design and implement a functional PC monitoring bot that combines RPA automation with web dashboard visualization. The scope includes backend monitoring, JSON file generation, Live Server dashboard hosting, frontend auto-refresh, debugging support, and data export.")
    page_break(doc)

    add_heading(doc, "CHAPTER 2 LITERATURE REVIEW")
    sections = [
        ("Introduction to System Monitoring", "System monitoring is the process of observing hardware and software resource usage in real time. It helps users understand how CPU, memory, GPU, and applications are performing. Tools like Windows Task Manager provide similar information, but custom dashboards allow project-specific visualization and automation."),
        ("Role of RPA in Monitoring", "Robotic Process Automation is commonly used to automate repetitive tasks, collect information from systems, and generate structured outputs. In this project, RPA is used to run monitoring logic repeatedly and write the result to a file that can be consumed by another application."),
        ("PowerShell for Performance Counters", "PowerShell provides access to Windows performance counters such as Processor Time, Memory Committed Bytes, and GPU Engine Utilization Percentage. These counters allow scripts to collect system usage data without requiring external monitoring software."),
        ("JSON-Based Data Exchange", "JSON is a lightweight data format used widely in web applications. It is human-readable, simple to generate from PowerShell, and easy to parse in JavaScript. This makes it suitable for communication between UiPath and the dashboard."),
        ("Web Dashboards for Real-Time Data", "HTML, CSS, and JavaScript can be used to create dashboards that update data dynamically. JavaScript fetch logic allows the page to request the latest JSON file repeatedly and update the interface without reloading the whole page."),
        ("Live Server in Development", "VS Code Live Server is useful for serving local HTML, CSS, JavaScript, and JSON files through HTTP. This is important because browser fetch requests work reliably when files are served through a local server instead of opened directly as file paths."),
        ("Exporting Monitoring Data", "Export options such as CSV and Excel files help users store monitoring snapshots for later analysis. Browser-based export logic can generate downloadable files directly from the currently loaded dashboard data."),
    ]
    for title, body in sections:
        add_heading(doc, title, 2)
        add_body(doc, body)
    page_break(doc)

    add_heading(doc, "CHAPTER 3")
    add_heading(doc, "METHODOLOGY AND SYSTEM IMPLEMENTATION")
    add_heading(doc, "EXISTING SYSTEM", 2)
    add_body(doc, "In the existing manual approach, users depend on Task Manager or Resource Monitor to check system performance. Although these tools are powerful, they do not provide a custom project dashboard, JSON data exchange, or automated export tailored to an RPA workflow.")
    add_heading(doc, "PROPOSED SYSTEM")
    add_body(doc, "The proposed system introduces a UiPath-based PC monitoring bot that collects system metrics automatically and writes them to a JSON file. A browser dashboard reads the file and displays live CPU, RAM, and GPU usage along with process-level application details.")
    add_heading(doc, "MODULES")
    add_body(doc, "The project is divided into backend monitoring, JSON generation, frontend dashboard rendering, auto-refresh logic, debugging support, and export functionality. Each module performs a specific role and together they form a complete monitoring system.")
    add_table(doc, ["Module", "Description"], [
        ("UiPath Backend", "Runs the monitoring workflow continuously and invokes the PowerShell script."),
        ("PowerShell Monitor", "Collects CPU, RAM, GPU, and application usage details from Windows."),
        ("JSON Writer", "Writes structured data to dashboard/performance.json using safe file replacement."),
        ("Dashboard Frontend", "Displays live metrics using HTML, CSS, and JavaScript."),
        ("Export Module", "Downloads the latest data as CSV or Excel-compatible XLS."),
    ])
    add_heading(doc, "OVERVIEW OF RPA")
    add_body(doc, "RPA allows software robots to perform repetitive digital tasks. In this project, the RPA bot repeatedly executes the monitoring script, ensuring that the dashboard receives fresh data without manual refresh or manual command execution.")
    add_heading(doc, "RPA FEATURES")
    add_body(doc, "The project demonstrates important RPA features such as repeated task execution, integration with PowerShell, file-based communication, error reduction, and automation of system data collection.")
    page_break(doc)

    add_heading(doc, "CHAPTER 4 RESULTS AND DISCUSSION")
    add_heading(doc, "RESULTS")
    add_body(doc, "The project successfully displays live CPU, RAM, and GPU values in the dashboard. It also displays the top applications using CPU and RAM. When GPU process data is available from Windows counters, the dashboard displays the applications using GPU along with the GPU engine and usage percentage.")
    add_table(doc, ["Output Area", "Result"], [
        ("CPU Card", "Shows total CPU usage with animated progress bar."),
        ("RAM Card", "Shows total committed memory usage with animated progress bar."),
        ("GPU Card", "Shows total GPU engine utilization."),
        ("Top CPU Apps", "Displays at least three CPU-consuming applications when available."),
        ("Top RAM Apps", "Displays at least three RAM-consuming applications when available."),
        ("Apps Using GPU", "Displays active GPU apps, or shows no active GPU apps when none are detected."),
        ("Export", "Downloads current data in CSV or Excel-compatible format."),
    ])
    paragraph(doc, "Fig 1: UiPath Main.xaml backend workflow")
    paragraph(doc, "Fig 2: PowerShell JSON generation")
    paragraph(doc, "Fig 3: Live dashboard with CPU, RAM, and GPU cards")
    paragraph(doc, "Fig 4: Top CPU and RAM applications table")
    paragraph(doc, "Fig 5: Applications using GPU table")
    paragraph(doc, "Fig 6: CSV and Excel export buttons")
    add_heading(doc, "FINAL OUTPUT")
    add_body(doc, "The final output is a fully functional dashboard that updates automatically every three seconds. The dashboard can be opened through VS Code Live Server and shows live system status using a clean dark interface.")
    add_heading(doc, "DISCUSSION")
    steps = [
        ("Step 1: Requirement Analysis", "The project requirement was to build a real-time dashboard similar to Task Manager using UiPath, HTML, CSS, JavaScript, and JSON file communication."),
        ("Step 2: Workflow Design", "The workflow was designed so UiPath continuously runs the backend monitor and updates the JSON file used by the frontend."),
        ("Step 3: Environment Setup", "The project folder was organized with dashboard files, logs, Main.xaml, MonitorPerformance.ps1, and project configuration."),
        ("Step 4: Backend Development", "PowerShell commands were used to collect CPU, RAM, GPU, and process usage data. The values were shaped into clean JSON fields."),
        ("Step 5: Frontend Development", "The dashboard was developed with dark UI styling, metric cards, progress bars, app tables, debug logs, and export buttons."),
        ("Step 6: JSON Fetch Fix", "The JSON file was placed inside the dashboard folder and fetched using ./performance.json to avoid Live Server parent-folder fetch issues."),
        ("Step 7: Testing and Debugging", "The project was tested by validating JSON output, checking browser fetch responses, and confirming that dashboard tables updated correctly."),
        ("Step 8: Export Implementation", "CSV and Excel-compatible export features were added so the current dashboard data can be downloaded for reporting."),
    ]
    for title, body in steps:
        add_heading(doc, title, 2)
        add_body(doc, body)
    page_break(doc)

    add_heading(doc, "CHAPTER 5 CONCLUSION AND FUTURE WORK")
    add_heading(doc, "CONCLUSION")
    add_body(doc, "The PC Performance Monitor Bot successfully demonstrates how UiPath can be used with PowerShell and a web dashboard to monitor live system performance. The project collects CPU, RAM, GPU, and application usage data and presents it in a modern browser interface.")
    add_body(doc, "The project also solves common beginner issues such as JSON fetch errors in Live Server, stale browser cache, and mismatched JSON field names. It provides a simple architecture that is easy to understand, test, and extend.")
    add_heading(doc, "FUTURE WORK")
    future = [
        "A future version can store historical monitoring data in a database instead of keeping only the latest JSON snapshot.",
        "Charts can be added to show CPU, RAM, and GPU usage trends over time.",
        "The bot can be extended to send alerts when CPU, RAM, or GPU usage crosses a defined threshold.",
        "Remote monitoring support can be added to observe multiple systems from a central dashboard.",
        "UiPath Orchestrator integration can be introduced for scheduled monitoring and centralized logging.",
        "The dashboard can include process filtering, search, and detailed drill-down views for advanced analysis.",
    ]
    for item in future:
        add_body(doc, item)
    page_break(doc)

    add_heading(doc, "REFERENCE")
    refs = [
        "https://www.uipath.com/rpa/robotic-process-automation - Learn how RPA automates repetitive business and system tasks.",
        "https://learn.microsoft.com/powershell/ - Microsoft PowerShell documentation.",
        "https://learn.microsoft.com/windows/win32/perfctrs/performance-counters-portal - Windows Performance Counters documentation.",
        "https://developer.mozilla.org/docs/Web/API/Fetch_API - JavaScript Fetch API documentation.",
        "https://developer.mozilla.org/docs/Learn/JavaScript/Objects/JSON - JSON format and parsing guide.",
        "https://marketplace.visualstudio.com/items?itemName=ritwickdey.LiveServer - VS Code Live Server extension.",
        "https://academy.uipath.com - UiPath learning resources.",
    ]
    for ref in refs:
        paragraph(doc, ref, style="List Paragraph", after=6)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
